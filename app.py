"""퇴사자 처리 자동화 앱 (로컬 실행).

실행:
    pip install -r requirements.txt
    playwright install chromium
    python app.py
    → 브라우저에서 http://127.0.0.1:5000 접속

흐름:
  1) 퇴사자 이름 + 퇴사일 입력
  2) 사이트 체크리스트가 순서대로 표시됨
  3) 각 사이트에서 [열기] → 사용자가 직접 로그인 → [자동 처리]
     - 브라우저 사이트: 저장/제출 직전에서 멈춤(사람이 최종 확인)
     - API 사이트(M365): 로그인 없이 바로 처리
  4) 처리 후 [완료] 로 체크. 진행상태는 저장돼 이어서 가능.
"""

from __future__ import annotations

# .env 를 먼저 읽어야 M365_*, BROWSER_CDP_URL 등이 반영된다.
# 실행 위치와 무관하게 app.py 옆의 .env 를 확실히 읽는다.
try:
    import os as _os

    from dotenv import load_dotenv

    load_dotenv(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), ".env"))
except Exception:
    pass

from flask import Flask, jsonify, render_template, request

from automation import registry, state
from automation.controller import controller
from automation.sites.base import Employee

app = Flask(__name__)

# 앱 시작 시: 이전 세션 작업은 보관함으로 옮기고 빈 화면으로 시작한다.
state.begin_session()

# 테스트 모드(실제 처리 없이 흐름만 확인). 앱 재시작하면 꺼진다.
_TEST_MODE = {"on": False}

# 처리에 포함할 사이트(기본 전체). 여기서 빠지면 순차 처리에서 제외된다.
_ENABLED_SITES = {s.id for s in registry.all_scenarios()}
# 앱 시작 후 메일 자동 불러오기를 1회만 하기 위한 플래그
_AUTO_IMPORTED = {"done": False}

# 메일 읽기 백그라운드 작업(진행상황 폴링용). 한 번에 하나만 돈다.
import threading  # noqa: E402

_MAIL_LOCK = threading.Lock()
_MAIL_JOB = {
    "running": False,
    "phase": "idle",   # idle | reading | done
    "fetched": 0,
    "cap": 0,
    "kind": None,      # import | list
    "result": None,
    "error": None,
}


def _mail_env():
    """.env 를 다시 읽어(즉시반영) 메일 조회 파라미터를 돌려준다."""
    import os

    try:
        from dotenv import load_dotenv

        load_dotenv(
            os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"),
            override=True,
        )
    except Exception:
        pass

    mailbox = os.environ.get("M365_MAILBOX", "").strip()
    # 핵심 키워드는 코드에 항상 포함(.py는 UTF-8이라 안 깨짐). .env 값은 추가로 합친다.
    keyword = (os.environ.get("MAIL_SUBJECT_KEYWORD", "") or "").strip() + ",퇴사,퇴직,퇴사자"
    # 발신자는 영문(이메일 주소)일 때만 필터로 쓴다. 한글 이름은 인코딩 문제로 무시.
    _s = os.environ.get("MAIL_SENDER", "").strip()
    sender = _s if (_s and _s.isascii()) else None
    try:
        since_days = int(os.environ.get("MAIL_SINCE_DAYS", "90") or "90")
    except ValueError:
        since_days = 90
    return mailbox, keyword, sender, since_days


def _date_block(resign_date: str) -> str | None:
    """퇴사일이 아직 안 됐으면 안내 문구를, 처리 가능하면 None 을 돌려준다."""
    import datetime as _dt

    try:
        rd = _dt.date.fromisoformat((resign_date or "").strip())
    except Exception:
        return None  # 날짜 형식이 이상하면 막지 않는다
    today = _dt.date.today()
    if today < rd:
        return f"퇴사일({resign_date}) 전입니다. {resign_date} 이후에 처리할 수 있습니다. (오늘 {today.isoformat()})"
    return None


def _do_mail_import(progress=None) -> dict:
    """메일을 읽어 새 퇴사자는 등록하고, 중복은 사용자 확인용으로 모아 돌려준다."""
    from automation import mailimport

    mailbox, keyword, sender, since_days = _mail_env()
    found, stats = mailimport.import_from_mail(
        mailbox, keyword, sender, since_days=since_days, progress=progress
    )

    # 이미 목록에 있는 사람(이름 기준)은 자동 추가하지 않고 사용자가 결정하도록 넘긴다.
    existing_names = {tg["name"] for tg in state.snapshot()["targets"]}
    site_ids = _site_ids()
    added, duplicates = [], []
    for t in found:
        if t["name"] in existing_names:
            duplicates.append(t)
        else:
            state.add_target(t["name"], t["resign_date"], site_ids)
            added.append(t)

    if found:
        parts = [f"메일({since_days}일 이내)에서 {len(added)}명 등록"]
        if duplicates:
            parts.append(f"중복 {len(duplicates)}명은 확인 필요")
        message = " · ".join(parts)
    else:
        # 진단: 어디서 걸렸는지 알려준다
        message = (
            f"등록 대상 없음 — 최근 {since_days}일 메일 {stats['total']}건 중 "
            f"제목매칭 {stats['subject_matched']}건, 이름·퇴사일 추출 {stats['parsed']}건."
        )
        if stats["subject_matched"] == 0:
            message += f" [키워드={stats.get('keywords')}]"
            toi = stats.get("subjects_with_toi") or []
            if toi:
                message += "\n'퇴' 들어간 제목: " + " | ".join(toi)
            if stats.get("recent_subjects"):
                message += "\n읽은 제목: " + " | ".join(stats["recent_subjects"])
        elif stats["matched_no_parse"]:
            message += " 파싱 실패 제목: " + ", ".join(stats["matched_no_parse"])

    return {
        "targets": state.snapshot(),
        "added": added,
        "duplicates": duplicates,
        "message": message,
    }


def _start_mail_job(kind: str) -> bool:
    """메일 읽기 백그라운드 작업을 시작한다. 이미 돌고 있으면 False."""
    with _MAIL_LOCK:
        if _MAIL_JOB["running"]:
            return False
        _MAIL_JOB.update(
            running=True, phase="reading", fetched=0, cap=0,
            kind=kind, result=None, error=None,
        )
    threading.Thread(target=_mail_worker, args=(kind,), daemon=True).start()
    return True


def _mail_configured() -> bool:
    """M365 메일 조회에 필요한 설정이 갖춰졌는지."""
    from automation.integrations.graph import GraphClient

    mailbox, *_ = _mail_env()
    try:
        return bool(GraphClient().configured() and mailbox)
    except Exception:
        return False


def _mail_progress(fetched: int, cap: int) -> None:
    _MAIL_JOB["fetched"] = fetched
    _MAIL_JOB["cap"] = cap


def _mail_worker(kind: str) -> None:
    """백그라운드에서 메일을 읽고 결과를 _MAIL_JOB 에 담는다(진행상황 폴링용)."""
    try:
        if kind == "list":
            from automation import mailimport

            mailbox, keyword, sender, since_days = _mail_env()
            rows = mailimport.debug_list(
                mailbox, keyword, sender, since_days=since_days, progress=_mail_progress
            )
            _MAIL_JOB["result"] = {"messages": rows}
        else:
            _MAIL_JOB["result"] = _do_mail_import(progress=_mail_progress)
    except Exception as exc:  # noqa: BLE001
        _MAIL_JOB["error"] = str(exc)
    finally:
        _MAIL_JOB["running"] = False
        _MAIL_JOB["phase"] = "done"


def _scenario_list() -> list[dict]:
    return [
        {
            "id": s.id,
            "name": s.name,
            "kind": s.kind,
            "url": s.url,
            "has_credentials": state.has_credentials(s.id),
        }
        for s in registry.all_scenarios()
    ]


@app.get("/")
def index():
    return render_template("index.html")


def _site_ids() -> list[str]:
    return [s.id for s in registry.all_scenarios()]


@app.get("/api/status")
def api_status():
    return jsonify(
        {
            "scenarios": _scenario_list(),
            "targets": state.snapshot(),
            "restorable": state.latest_restorable(),
            "test_mode": _TEST_MODE["on"],
            "enabled_sites": sorted(_ENABLED_SITES),
        }
    )


@app.post("/api/sites/enabled")
def api_sites_enabled():
    ids = (request.get_json(force=True, silent=True) or {}).get("ids")
    if isinstance(ids, list):
        _ENABLED_SITES.clear()
        _ENABLED_SITES.update(str(i) for i in ids)
    return jsonify({"enabled_sites": sorted(_ENABLED_SITES)})


@app.post("/api/mail/auto-start")
def api_mail_auto_start():
    """앱 시작 후 1회, 백그라운드로 자동 메일 가져오기를 시작한다.

    설정 안 됐거나 이미 한 번 했으면 skip. 시작하면 클라이언트가
    /api/mail/progress 로 진행바를 띄우고 결과를 받는다.
    """
    if _AUTO_IMPORTED["done"]:
        return jsonify({"skip": True, "reason": "already"})
    if not _mail_configured():
        _AUTO_IMPORTED["done"] = True
        return jsonify({"skip": True, "reason": "not_configured"})
    _AUTO_IMPORTED["done"] = True
    _start_mail_job("import")
    return jsonify({"started": True})


@app.post("/api/testmode")
def api_testmode():
    data = request.get_json(force=True, silent=True) or {}
    _TEST_MODE["on"] = bool(data.get("on"))
    return jsonify({"test_mode": _TEST_MODE["on"]})


@app.post("/api/site/<site_id>/credentials")
def api_set_credentials(site_id: str):
    if registry.get(site_id) is None:
        return jsonify({"error": "알 수 없는 사이트"}), 404
    data = request.get_json(force=True, silent=True) or {}
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""
    if not username:
        state.clear_credentials(site_id)  # 아이디 비우면 저장 해제
    else:
        state.set_credentials(site_id, username, password)
    return jsonify({"has_credentials": state.has_credentials(site_id)})


@app.get("/api/history")
def api_history():
    return jsonify({"history": state.list_history()})


@app.post("/api/history/restore")
def api_history_restore():
    session_id = (request.get_json(force=True, silent=True) or {}).get("id")
    state.restore(session_id)
    return jsonify({"targets": state.snapshot()})


@app.post("/api/target")
def api_target():
    """대상자 1명 추가(수동). 추가한 사람을 활성 대상으로 만든다."""
    data = request.get_json(force=True)
    name = (data.get("name") or "").strip()
    resign_date = (data.get("resign_date") or "").strip()
    if not name or not resign_date:
        return jsonify({"error": "이름과 퇴사일을 모두 입력하세요."}), 400
    state.add_target(name, resign_date, _site_ids(), make_active=True)
    return jsonify({"targets": state.snapshot()})


@app.post("/api/target/active")
def api_target_active():
    key = (request.get_json(force=True).get("key") or "").strip()
    if not state.set_active(key):
        return jsonify({"error": "대상을 찾지 못했습니다."}), 404
    return jsonify({"targets": state.snapshot()})


@app.post("/api/target/remove")
def api_target_remove():
    key = (request.get_json(force=True).get("key") or "").strip()
    state.remove_target(key)
    return jsonify({"targets": state.snapshot()})


@app.post("/api/mail/start")
def api_mail_start():
    """메일 읽기를 백그라운드로 시작한다. 진행상황은 /api/mail/progress 로 폴링."""
    kind = (request.get_json(force=True, silent=True) or {}).get("kind") or "import"
    if kind not in ("import", "list"):
        kind = "import"
    if not _start_mail_job(kind):
        return jsonify({"running": True, "already": True, "kind": _MAIL_JOB["kind"]})
    return jsonify({"started": True, "kind": kind})


@app.get("/api/mail/progress")
def api_mail_progress():
    """메일 읽기 진행상황(퍼센트)과, 끝났으면 결과를 함께 돌려준다."""
    j = _MAIL_JOB
    if j["cap"]:
        pct = min(99, int(j["fetched"] * 100 / j["cap"]))
    else:
        pct = 5 if j["running"] else 0
    done = (not j["running"]) and j["phase"] == "done"
    if done and not j["error"]:
        pct = 100
    return jsonify(
        {
            "running": j["running"],
            "phase": j["phase"],
            "fetched": j["fetched"],
            "cap": j["cap"],
            "percent": pct,
            "kind": j["kind"],
            "error": j["error"],
            "result": j["result"] if done else None,
        }
    )


@app.post("/api/target/reset")
def api_target_reset():
    key = (request.get_json(force=True).get("key") or "").strip()
    state.reset_target(key)
    return jsonify({"targets": state.snapshot()})


@app.post("/api/target/done")
def api_target_done_all():
    """이 대상자(사람)의 모든 시스템을 완료로 표시한다(완료는 사람에게 귀속)."""
    key = (request.get_json(force=True).get("key") or "").strip()
    tg = next((t for t in state.snapshot()["targets"] if t["key"] == key), None)
    if not tg:
        return jsonify({"error": "대상을 찾지 못했습니다."}), 404
    for sid in _site_ids():
        state.set_site(tg["name"], tg["resign_date"], sid, "done", "완료 처리됨")
    return jsonify({"targets": state.snapshot()})


@app.post("/api/targets/clear")
def api_targets_clear():
    state.clear_all()
    return jsonify({"targets": state.snapshot()})


def _fmt_date(v) -> str:
    """엑셀 셀 값을 YYYY-MM-DD 문자열로 정규화한다."""
    import datetime

    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    digits = "".join(ch for ch in s if ch.isdigit())
    if len(digits) == 8:  # 20260715 → 2026-07-15
        return f"{digits[0:4]}-{digits[4:6]}-{digits[6:8]}"
    return s.replace(".", "-").replace("/", "-").replace(" ", "")


def _parse_excel(file_storage) -> list[dict]:
    """이름/퇴사일 열이 있는 엑셀을 읽어 대상자 목록을 만든다.

    헤더에 '이름/성명', '퇴사/퇴직/날짜/일자' 가 있으면 그 열을 쓰고,
    없으면 첫 열=이름, 둘째 열=퇴사일 로 간주한다.
    """
    import io

    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(file_storage.read()), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    def _txt(v) -> str:
        return str(v).strip() if v is not None else ""

    name_idx, date_idx, start = 0, 1, 0
    header = [_txt(c) for c in rows[0]]
    found = False
    for i, t in enumerate(header):
        if any(k in t for k in ("이름", "성명")):
            name_idx, found = i, True
        if any(k in t for k in ("퇴사", "퇴직", "날짜", "일자")):
            date_idx, found = i, True
    if found:
        start = 1  # 헤더 행은 건너뛴다

    targets = []
    for row in rows[start:]:
        if not row:
            continue
        name = _txt(row[name_idx]) if name_idx < len(row) else ""
        raw = row[date_idx] if date_idx < len(row) else None
        if not name:
            continue
        targets.append({"name": name, "resign_date": _fmt_date(raw)})
    return targets


@app.post("/api/target/excel")
def api_target_excel():
    f = request.files.get("file")
    if f is None or not f.filename:
        return jsonify({"error": "엑셀 파일을 선택하세요."}), 400
    try:
        targets = _parse_excel(f)
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": f"엑셀을 읽지 못했습니다: {exc}"}), 400
    if not targets:
        return jsonify({"error": "대상자를 찾지 못했습니다. (이름/퇴사일 열을 확인해 주세요)"}), 400
    # 파싱한 대상자 전원을 목록에 추가한다(활성은 기존 유지 or 첫 대상).
    site_ids = _site_ids()
    for t in targets:
        if t.get("resign_date"):
            state.add_target(t["name"], t["resign_date"], site_ids)
    return jsonify({"targets": state.snapshot(), "added": len(targets)})


def _current_employee() -> Employee | None:
    cur = state.active_entry()
    if not cur:
        return None
    return Employee(name=cur["name"], resign_date=cur["resign_date"])


@app.post("/api/site/<site_id>/open")
def api_open(site_id: str):
    scenario = registry.get(site_id)
    emp = _current_employee()
    if scenario is None:
        return jsonify({"error": "알 수 없는 사이트"}), 404
    if emp is None:
        return jsonify({"error": "먼저 퇴사자를 입력하세요."}), 400
    if scenario.kind != "browser":
        return jsonify({"error": "이 사이트는 브라우저를 열지 않습니다(API 처리)."}), 400
    try:
        controller.open_site(scenario.url)
        state.set_site(emp.name, emp.resign_date, site_id, "in_progress", "로그인 대기 중")
        return jsonify({"ok": True, "message": "브라우저에서 로그인 후 [자동 처리]를 누르세요."})
    except Exception as exc:  # noqa: BLE001
        state.set_site(emp.name, emp.resign_date, site_id, "error", str(exc))
        return jsonify({"error": str(exc)}), 500


@app.post("/api/site/<site_id>/run")
def api_run(site_id: str):
    scenario = registry.get(site_id)
    emp = _current_employee()
    if scenario is None:
        return jsonify({"error": "알 수 없는 사이트"}), 404
    if emp is None:
        return jsonify({"error": "먼저 퇴사자를 입력하세요."}), 400
    # 퇴사일 이전이면 처리 불가(테스트 모드는 예외로 통과시켜 흐름 확인 가능)
    block = _date_block(emp.resign_date)
    if block and not _TEST_MODE["on"]:
        state.set_site(emp.name, emp.resign_date, site_id, "blocked", block)
        return jsonify({"ok": True, "awaiting": False, "blocked": True, "message": block})
    try:
        scenario.test_mode = _TEST_MODE["on"]  # 실행 직전 테스트 모드 반영
        scenario.credentials = state.get_credentials(site_id)
        if scenario.kind == "api":
            result = scenario.run_api(emp)
        else:
            result = controller.run_scenario(scenario, emp)

        if not result.ok:
            status = "error"
        elif result.awaiting:
            status = "awaiting"
        else:
            status = "done"
        state.set_site(emp.name, emp.resign_date, site_id, status, result.message)
        return jsonify(
            {"ok": result.ok, "awaiting": result.awaiting, "message": result.message}
        )
    except Exception as exc:  # noqa: BLE001
        state.set_site(emp.name, emp.resign_date, site_id, "error", str(exc))
        return jsonify({"error": str(exc)}), 500


@app.post("/api/run-all")
def api_run_all():
    """목록의 모든 대상자를 순서대로 처리한다(브라우저는 미리 [열기]+로그인 필요).

    - 테스트 모드: 각자 팝업까지 확인 후 닫고 다음 사람으로 (끝까지 자동).
    - 실제 모드: 사람이 [완료]해야 하는 지점(awaiting)에서 멈추고 안내한다.
    """
    body = request.get_json(force=True, silent=True) or {}
    only_keys = body.get("keys")  # 선택한 대상만 처리(없으면 전체)
    # 포함(체크된) 사이트만 순서대로
    order = [s for s in registry.all_scenarios() if s.id in _ENABLED_SITES]
    processed = 0
    errors = []
    opened = set()  # 사이트별로 브라우저는 한 번만 연다(사람마다 재이동 방지)
    for tg in state.snapshot()["targets"]:
        if only_keys and tg["key"] not in only_keys:
            continue
        emp = Employee(name=tg["name"], resign_date=tg["resign_date"])
        # 퇴사일 이전이면 이 사람은 건너뛰고 blocked 표시(테스트 모드는 예외)
        block = _date_block(emp.resign_date)
        if block and not _TEST_MODE["on"]:
            for s in order:
                if tg["sites"].get(s.id, {}).get("status") != "done":
                    state.set_site(emp.name, emp.resign_date, s.id, "blocked", block)
            continue
        for s in order:
            if tg["sites"].get(s.id, {}).get("status") == "done":
                continue
            try:
                s.test_mode = _TEST_MODE["on"]
                s.credentials = state.get_credentials(s.id)
                if s.kind == "api":
                    result = s.run_api(emp)
                else:
                    if s.url not in opened:
                        controller.open_site(s.url)  # 최초 1회만 (이후엔 바로 검색)
                        opened.add(s.url)
                    result = controller.run_scenario(s, emp)
            except Exception as exc:  # noqa: BLE001
                state.set_site(emp.name, emp.resign_date, s.id, "error", str(exc))
                errors.append(f"{emp.name}·{s.name}")
                continue

            status = "error" if not result.ok else ("awaiting" if result.awaiting else "done")
            state.set_site(emp.name, emp.resign_date, s.id, status, result.message)
            processed += 1
            if status == "awaiting":
                return jsonify(
                    {
                        "targets": state.snapshot(),
                        "message": (
                            f"{emp.name} · {s.name}: 실제 퇴사라서 [완료] 직전에 멈췄습니다. "
                            "브라우저에서 완료한 뒤 다시 [전체 순차 처리]로 이어가세요.\n\n"
                            "※ 실제 처리 없이 여러 명 흐름만 테스트하려면 위의 🧪 테스트 모드를 켜세요 "
                            "(팝업까지 확인 후 자동으로 닫고 다음 사람으로 넘어갑니다)."
                        ),
                    }
                )
            if status == "error":
                errors.append(f"{emp.name}·{s.name}")

    # 사람별 결과를 그대로 보여준다(어디까지 됐는지 투명하게).
    label = {
        "done": "완료", "awaiting": "완료대기", "error": "실패",
        "pending": "대기", "blocked": "퇴사일전",
    }
    final = state.snapshot()["targets"]
    lines = []
    for tg in final:
        if only_keys and tg["key"] not in only_keys:
            continue
        parts = []
        for s in order:
            st = tg["sites"].get(s.id, {}).get("status", "pending")
            parts.append(f"{s.name}={label.get(st, st)}")
        lines.append(f"· {tg['name']}: " + ", ".join(parts))
    msg = f"순차 처리 종료 ({len(lines)}명)\n" + "\n".join(lines)
    return jsonify({"targets": state.snapshot(), "message": msg})


@app.post("/api/site/<site_id>/done")
def api_done(site_id: str):
    emp = _current_employee()
    if emp is None:
        return jsonify({"error": "먼저 퇴사자를 입력하세요."}), 400
    state.set_site(emp.name, emp.resign_date, site_id, "done", "완료 처리됨")
    return jsonify({"ok": True})


if __name__ == "__main__":
    # 로컬 전용. 회사 PC 한 대에서 사용하는 것을 전제로 한다.
    app.run(host="127.0.0.1", port=5000, debug=False)
